import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_presentation(output_path="Python_Data_Analysis_and_Deployment.pptx"):
    prs = Presentation()
    # 设置 16:9 现代宽屏尺寸
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]  # 空白版式

    # 配色规范（极简白底、高对比度阅读）
    COLOR_BG = RGBColor(255, 255, 255)          # 纯白背景
    COLOR_TITLE = RGBColor(15, 23, 42)          # 深石板灰标题
    COLOR_SUBTITLE = RGBColor(99, 102, 241)     # 靛青模块标签
    COLOR_TEXT = RGBColor(51, 65, 85)           # 正文深灰（易读、不刺眼）
    COLOR_BOX_BG = RGBColor(248, 250, 252)      # 代码框浅灰底
    COLOR_BOX_BORDER = RGBColor(226, 232, 240)  # 代码框浅色边框
    COLOR_CODE = RGBColor(30, 41, 59)           # 代码深色
    COLOR_CALLOUT_BG = RGBColor(241, 245, 249)  # 金句框微蓝灰底
    COLOR_CALLOUT_BORDER = RGBColor(199, 210, 254)
    COLOR_CALLOUT_TEXT = RGBColor(30, 27, 75)

    FONT_MAIN = "Microsoft YaHei"
    FONT_CODE = "Consolas"

    slides_data = [
        # --- 模块一：思维跃迁 ---
        {
            "kicker": "模块一 · 思维跃迁",
            "title": "从工程后端思维回归数据本质",
            "bullets": [
                "后端工程关注对象封装、网络 I/O、状态流转与线程锁安全；而数据科学则转向数据导向设计（Data-Oriented Design），追求数据与操作彻底解耦。",
                "原生 Python 的 for 循环受制于解释器逐行解析与动态类型推断；现代矩阵思维依赖连续物理内存排布与底层 SIMD 向量化并发吞吐。",
                "Python 统治数据科学与 AI 的本质是其无与伦比的 C-API 扩展胶水能力：上层动态交互，下层直驱 OpenBLAS、MKL 与 CUDA 底层库。",
                "本课程目标：构建从底层数组原理（NumPy）到表格分析（Pandas），再到文本模型微调（MiniRBT）与轻量化部署（ONNX）的完整工程闭环。"
            ],
            "code": "# 向量化思维：单指令多数据并发吞吐\nresult = prices * 1.05 + 3.0",
            "takeaway": "核心金句：学 NumPy 和 Pandas 不是记 API，而是彻底丢弃逐个遍历的直觉，学会用整块连续数据去思考。"
        },
        {
            "kicker": "模块一 · 思维跃迁",
            "title": "原生 Python 处理海量数据的瓶颈",
            "bullets": [
                "指针数组与内存碎片化：Python list 存放的是离散对象的内存指针（每个整型 PyObject 占 28 字节以上），无法利用现代 CPU 的缓存预取（Cache Miss 频发）。",
                "解释器动态调度开销：每一轮原生 for 循环，解释器都要付出对象遍历、类型核对、引用计数变更与字节码跳转开销，百万级循环累积显著耗时。",
                "NumPy 将同质数据（相同 dtype）分配在一整段平坦、连续的物理缓冲区（Buffer）；Pandas 在此基础上增添了标签轴（Index）与关系代数语义。",
                "结论：原生 Python 适合处理灵活异构对象，但在百万级同构数值计算中，NumPy 和 Pandas 的吞吐效率呈数量级优势。"
            ],
            "code": "# 连续物理内存分配，直抵底层 C 求和\narr = np.arange(1_000_000, dtype=np.int64)\ntotal = arr.sum()",
            "takeaway": "核心金句：原生 Python 能做，但在大批量数据上极不划算；专业工具的价值是让批量处理更贴近底层硬件。"
        },
        {
            "kicker": "模块一 · 思维跃迁",
            "title": "分析载体：NYC 出租车数据与 Parquet 革命",
            "bullets": [
                "工业级异构样本：纽约市 TLC 黄色出租车数据包含供应商、上下车精确时间戳、运距、费率、分区 ID 及明细费用项等 20 个混合字段。",
                "行式纯文本（CSV）的局限：无 Schema 约束，读取时必须全表字符流扫描并逐行切分字符串推断类型，I/O 冗余高且解析极慢。",
                "列式存储（Parquet）的优势：内嵌强类型 Schema 与字典编码，支持高倍率压缩，并原生支持列裁剪（仅读取指定列）与分区跳读下推。",
                "全课业务五问贯穿体系：围绕数据可信度校验（Q1）、时域规律（Q2）、运营效率（Q3）、空间拓扑（Q4）与窗口异常（Q5）层层递进。"
            ],
            "code": "# 高性能读取 Parquet，自动映射强类型与时间戳\ntaxi = pd.read_parquet('../data/yellow_tripdata_2025-12.parquet')",
            "takeaway": "核心金句：大数据分析的第一步在 I/O；Parquet 列式存储在物理底层就为后续高性能切片扫清了障碍。"
        },

        # --- 模块二：NumPy 底层解密 ---
        {
            "kicker": "模块二 · NumPy 底层解密",
            "title": "解剖 ndarray：内存块与元数据驱动",
            "bullets": [
                "ndarray 并非更快的 list，在底层 C 结构体中它由 Data Buffer（数据缓冲区）与 Metadata（元数据信息）严格分离构成。",
                "dtype（数据类型）：规定单元素物理字节宽度与编码方式（如 int32 占 4 字节，float64 占 8 字节）。",
                "shape（逻辑形状）：赋予一维连续字节流多维的业务解释视角，如 (3, 4) 对应 3 行 4 列的逻辑矩阵。",
                "strides（步长元组）：定义逻辑坐标沿各维度移动一步时，指针在底层平坦物理内存中需跨越的实际字节数（Offset = i*s0 + j*s1）。"
            ],
            "code": "arr = np.array([[1, 2], [3, 4]], dtype=np.int32)\nprint(arr.shape)    # 逻辑形状: (2, 2)\nprint(arr.strides)  # 物理步长: (8, 4) 字节",
            "takeaway": "核心金句：数据物理上永远是一维线性的，多维矩阵只是元数据给予我们的一种数学视角。"
        },
        {
            "kicker": "模块二 · NumPy 底层解密",
            "title": "内存机制：View（视图）、Copy 与 Strides",
            "bullets": [
                "零拷贝视图（View）：基础切片（arr[0:2]）与矩阵转置（arr.T）仅生成新元数据，修改 shape 与 strides，物理数据零复制，耗时 O(1)。",
                "View 的副作用：因共享物理底表，对视图元素的写操作会同步影响原数组；若需环境隔离必须显式调用 .copy() 开辟全新内存。",
                "花式索引（Fancy Indexing）必深拷贝：由于布尔掩码（arr > 0）与离散索引抽取的元素在物理地址上非等差分布，无法用统一步长映射。",
                "连续性陷阱：转置后的数组破坏了行优先连续性（C-Contiguous），此时强行 reshape 可能会隐式触发数据搬运与内存重排。"
            ],
            "code": "sub_view = arr[0:2, :]  # 共享内存的 View\nsafe_copy = arr.copy()  # 独立物理深拷贝",
            "takeaway": "核心金句：NumPy 之所以快，很大程度在于它不急着复制数据；切片只是调整观察窗口，copy 才是真正搬砖。"
        },
        {
            "kicker": "模块二 · NumPy 底层解密",
            "title": "张量维度导航：Axis 轴的本质",
            "bullets": [
                "几何混淆破除：初学者常误将 axis=0 理解为“计算每一行”，导致方向判断完全倒置。",
                "降维核心法则：指定哪个 axis，就是“沿着该维度的索引变化方向穿透压缩”，该维度将在结果形状中直接消失。",
                "二维物理行为：axis=0 沿垂直跨行方向坍缩，消除第 0 维输出各列汇总（列和）；axis=1 沿水平跨列方向压缩，输出各行汇总（行和）。",
                "高维多轴压缩：在 (Batch, Channel, Height, Width) 的深度学习张量中，axis=(2, 3) 表示同时抹平空间高宽，保留通道特征。"
            ],
            "code": "# axis=0: 跨行纵向压缩 -> 输出列求和\ncol_sum = matrix.sum(axis=0)\n# axis=1: 跨列横向压缩 -> 输出行求和\nrow_sum = matrix.sum(axis=1)",
            "takeaway": "核心金句：记住这个口诀——“指定哪轴，哪轴消失”；轴操作的本质是高维向低维的投影坍缩。"
        },
        {
            "kicker": "模块二 · NumPy 底层解密",
            "title": "广播机制（Broadcasting）：虚拟维度扩展",
            "bullets": [
                "对齐规则：从两个数组 shape 的最右侧（尾部维度）向前对齐比对，若对应维度大小严格相等、或其中一个大小为 1，则允许广播。",
                "步长为 0 的秘密：当 (1, N) 向量被广播扩展为 (M, N) 时，NumPy 绝不真正复制数据，而是将其第 0 轴步长设为 0（strides[0]=0）。",
                "物理意义：跨行寻址时指针物理偏移量为 0，反复读取同一段内存，实现零额外内存开销的逻辑扩展。",
                "维度升维（newaxis）：利用 np.newaxis 可以在指定位置插入长度为 1 的轴，瞬间完成多维网格组合与样本成对距离计算。"
            ],
            "code": "# 利用广播特性，零循环计算两两欧氏距离\ndiff = pts[:, np.newaxis, :] - pts[np.newaxis, :, :]\ndist = np.sqrt(np.sum(diff**2, axis=-1))",
            "takeaway": "核心金句：广播机制不是在内存里把数据复制一万遍，而是巧妙地把指针的移动步长设成了 0。"
        },
        {
            "kicker": "模块二 · NumPy 底层解密",
            "title": "向量化（Vectorization）思维的升维",
            "bullets": [
                "思想内核：向量化并不是发明新算法，而是把逐个元素的指令循环，改写为作用于整块数组的声明式数学表达式。",
                "底层 SIMD 硬件赋能：编译后的底层 C 循环可直接映射到 CPU 向量寄存器（AVX-512、Neon），单条硬件指令并发处理多个浮点数。",
                "布尔掩码（Boolean Mask）：以批处理替代 if-else 业务分支，一次性生成整列 True/False，彻底消除 CPU 分支预测失败开销。",
                "编写建议：能直接写成整列或整块矩阵表达式时，绝不手写 Python 原生 for 循环，兼顾代码高雅性与运行吞吐量。"
            ],
            "code": "# 声明式布尔掩码批量过滤与上浮\nmask = trip_distance > 0\nvalid_prices = prices[mask] * 1.05 + 3.0",
            "takeaway": "核心金句：在数据科学里，代码中每多写一个 Python 原生 for 循环，底层向量计算单元都在哭泣一次。"
        },

        # --- 模块三：Pandas 核心基石 ---
        {
            "kicker": "模块三 · Pandas 核心基石",
            "title": "从纯矩阵到业务实体：Series、DF 与 Index",
            "bullets": [
                "业务痛点：ndarray 擅长纯数值密集计算，但无法承载真实业务表中的列名、字符串、时间戳、缺失值（NaN）与多表键关联。",
                "Series 的本质：一维带标签数组。不仅包含底层数值缓冲，还拥有字段名与不可变的轴标签（Index）。",
                "DataFrame 的本质：共享行索引（Index）的二维异构表格容器，逻辑上是关系表，物理上是按列组织的一组 Series。",
                "标签自动对齐（Alignment）：两表或两列相加时，Pandas 自动按 Index 键匹配相加，无法对齐的位置安全置为 NaN，绝不混淆行顺序。"
            ],
            "code": "s1 = pd.Series([10, 20], index=['A', 'B'])\ns2 = pd.Series([1, 2], index=['B', 'C'])\n# 自动按 Index 标签对齐求和，A与C为 NaN\nres = s1 + s2",
            "takeaway": "核心金句：NumPy 负责底层算得快，Pandas 负责让人看得懂；Index 是数据的身份证，只要身份证对齐，数据就不会乱。"
        },
        {
            "kicker": "模块三 · Pandas 核心基石",
            "title": "深入 Index：从业务主键到时间窗口",
            "bullets": [
                "Index 绝对不只是自增行号，它是用于唯一标识、精准寻址、哈希检索与关系代数对齐的核心锚点。",
                "不可变性（Immutable）：Index 对象创建后无法随意单点篡改，保证了链式数据流转与多表操作中的引用一致性与哈希安全。",
                "形态演进：从极省内存的虚拟 RangeIndex，到业务主键字符串 Index，再到高阶的 MultiIndex（层次化索引）。",
                "DatetimeIndex 原生时序能力：将时间字段置为 Index 后，可直接使用人类可读的字符串进行范围切片与周期重采样。"
            ],
            "code": "df = df.set_index('pickup_datetime')\n# 原生语义级时序范围切片\nmorning_orders = df.loc['2025-12-01 08:00':'2025-12-01 09:00']",
            "takeaway": "核心金句：永远别只把 Index 当行号；把时间戳设为 Index，整个 DataFrame 就瞬间变身为一台时序分析引擎。"
        },

        # --- 模块四：业务实战五问 ---
        {
            "kicker": "模块四 · 业务实战五问",
            "title": "实战 Q1：这 200 万条数据真的可靠吗？",
            "bullets": [
                "垃圾进，垃圾出（Garbage in, Garbage out）：任何未经验证的真实数据都存在脏值、类型偏移、计价断电、负数退费等干扰。",
                "三维诊断闭环：类型核对（dtypes 确认时间是否为 object）、空值率核查（isna().sum()）、物理业务边界防线搭建。",
                "多重业务守卫：基础车费 fare > 0、总金额 total > 0、行程距离 dist > 0、时长 minutes > 0、分区 ID 合规。",
                "细粒度错误归因（np.select）：不仅做布尔过滤，更要使用 np.select 对异常行打上 bad_fare、bad_duration 等标签形成审计看板。"
            ],
            "code": "# 批量规则打标与可信数据快速准入\nconds = [taxi['fare_amount'] <= 0, taxi['trip_distance'] <= 0]\nlabels = ['bad_fare', 'bad_distance']\ntaxi['flag'] = np.select(conds, labels, default='ok')\nclean_df = taxi.query(\"flag == 'ok'\").copy()",
            "takeaway": "核心金句：数据分析的第一要义不是建模，而是用怀疑的眼光严密体检数据质量边界。"
        },
        {
            "kicker": "模块四 · 业务实战五问",
            "title": "实战 Q2：一天中什么时候最忙？",
            "bullets": [
                "时域特征派生：出租车精确时间戳无法直接反映趋势，需利用 .dt.hour 与 .dt.day_name() 提取小时与工作日维度。",
                "GroupBy 机制（Split-Apply-Combine）：底层先按键哈希分桶切片，组内并行执行聚合函数，最后拼装为规整报表。",
                "现代命名聚合规范（Named Aggregation）：采用 agg(new_col=('col', 'func'))，一次扫描产出订单数、总流水与均价。",
                "多维联合分组：通过 ['pickup_weekday', 'pickup_hour'] 联合分组，清晰解构工作日通勤尖峰与周末深夜娱乐出行模式。"
            ],
            "code": "hourly = clean_df.assign(hour=clean_df['tpep_pickup_datetime'].dt.hour) \\\n    .groupby('hour', as_index=False) \\\n    .agg(orders=('total_amount', 'size'),\n         revenue=('total_amount', 'sum'))",
            "takeaway": "核心金句：聚合统计不是为了把数据做小，而是通过维度切片把混杂的 200 万条记录沉淀为清晰的决策报表。"
        },
        {
            "kicker": "模块四 · 业务实战五问",
            "title": "实战 Q3：哪个时段流水与运营效率最高？",
            "bullets": [
                "告别单一绝对值：总流水高可能仅因车多或拥堵；精细化运营必须关注单位时间产出（Revenue per Minute, RPM）。",
                "构造核心派生效率指标：每分钟收益（total_amount / trip_minutes）与每英里费率（fare_amount / trip_distance）。",
                "Transform 的核心价值：与 agg 压缩行数不同，transform 在组内计算统计量后，将结果如广播般回填到原表每一行中。",
                "相对偏离度分析：每笔订单 RPM 除以该小时组内均值（hour_avg_rpm），迅速揪出不受路况影响的高创收神单。"
            ],
            "code": "# 组内计算均值并回填每一行，不改变表规模\nq3 = clean_df.assign(rpm=clean_df['total_amount'] / clean_df['trip_minutes'])\nq3['hour_avg_rpm'] = q3.groupby('pickup_hour')['rpm'].transform('mean')\nq3['efficiency_ratio'] = q3['rpm'] / q3['hour_avg_rpm']",
            "takeaway": "核心金句：想看大盘概貌用 agg；想看个体在所属圈子里处于什么相对水平，必须用 transform。"
        },
        {
            "kicker": "模块四 · 业务实战五问",
            "title": "实战 Q4：哪些上下车区域最热门？",
            "bullets": [
                "双轴频次统计：出租车记录包含起点（PULocationID）与终点（DOLocationID），分别利用 value_counts() 获得离散分布。",
                "外连接（Outer Join）防丢失：部分偏远区域可能在特定月份“只有流出没有流入”，全外连接配合 fillna(0) 保障空间完整性。",
                "连接防御性校验（validate）：利用 validate='one_to_one' 预防因主键重复引发的笛卡尔积爆炸（行数翻倍致命 Bug）。",
                "关联官方地理维表：使用 validate='many_to_one' 将纯数字 LocationID 映射为 Manhattan、JFK 机场等真实行政区与地名。"
            ],
            "code": "# 频次汇总与带严密校验的双表关联\nhotspots = pu_cnt.merge(do_cnt, on='LocationID', how='outer', validate='1:1') \\\n    .fillna(0) \\\n    .assign(total=lambda x: x['pu'] + x['do']) \\\n    .merge(zone_dim, on='LocationID', how='left', validate='m:1')",
            "takeaway": "核心金句：连表不加 validate 等于开车不系安全带；别让脏数据制造的隐蔽笛卡尔积毁掉你的整张底表。"
        },
        {
            "kicker": "模块四 · 业务实战五问",
            "title": "实战 Q5：哪些订单和时间窗口可能异常？",
            "bullets": [
                "窗口思维（Window Thinking）：不能孤立判断数值大小；暴风雪夜的长途单 80 美元合情合理，但在局部小时内可能就是离群点。",
                "微观分位数分类：采用 quantile(0.01) 与 quantile(0.99) 截断极端值，结合 where 算子快速抽查异常订单。",
                "宏观时序分桶（Resample）：相当于专属于时间轴的 GroupBy，按 '1h' 网格化将连续离散事件重采样为平稳时间序列。",
                "动态 3-Sigma 滚动阈值（Rolling）：计算过去 6 小时的动态滑动均值与标准差，构建随趋势波动的上界，捕获流水真实突增尖峰。"
            ],
            "code": "# Resample 时间分桶与 Rolling 动态阈值构建\nts = clean_df.set_index('pickup_datetime')['total_amount'].resample('1h').sum()\nts_roll = ts.rolling(6, min_periods=3)\nupper_bound = ts_roll.mean() + 3 * ts_roll.std()\nspikes = ts[ts > upper_bound]",
            "takeaway": "核心金句：静态写死阈值只能应对死规则，动态滑动窗口才能看清活业务；3-Sigma 随业务节奏一起呼吸。"
        },

        # --- 模块五：文本分析与模型微调 ---
        {
            "kicker": "模块五 · 文本分析与小模型微调",
            "title": "非结构化跃迁：客服意图识别任务",
            "bullets": [
                "形态跃迁：真实业务中不仅有表格数值，更存在海量非结构化文本（用户咨询、工单描述、客服对话与搜索意图）。",
                "意图识别（Intent Classification）：自然语言处理的基石任务，输入一段自由文本，判别其归属的业务场景标签并分流路由。",
                "真实教学语料画像：1,599 条 JSONL 样本，涵盖四大金融意图类别：还款（696）、收支分析（399）、申卡（332）、优惠（172）。",
                "文本长度探索：P50 长度 43 字，P95 仅 63 字；明确了 max_length=64 即能覆盖绝大多数文本语义，无需长文本切块。"
            ],
            "code": "# 文本数据同样遵循先转为 DataFrame 分析的准则\nintent_df = pd.DataFrame(json_records)\nprint(intent_df['label'].value_counts())\nprint(intent_df['text'].str.len().describe())",
            "takeaway": "核心金句：自然语言进模型前依然是一张 DataFrame；连文本长度分布都没看过就去调超参数，属于盲人摸象。"
        },
        {
            "kicker": "模块五 · 文本分析与小模型微调",
            "title": "开发环境：Conda 沙箱与最小依赖链",
            "bullets": [
                "深度学习环境痛点：底层 CUDA 驱动、C++ 编译运行时与 Python 包版本容易冲突，必须进行系统级严格隔离。",
                "venv vs Conda：venv 仅隔离 Python 纯 Wheel 包；Conda 能独立封装特定版本 Python 解释器二进制与底层 C 动态库，稳定性更强。",
                "最小化工具链哲学：刻意剥离 DeepSpeed、bitsandbytes 等繁复重型框架，聚焦极简可复现的核心依赖链。",
                "跨硬件平台兼容：本套最小依赖支持 Linux、Windows CPU 以及 Apple Silicon 芯片（MPS 加速），保证教学全员跑通。"
            ],
            "code": "# 纯净轻量环境准备\nconda create -n minirbt_intent python=3.10 -y\nconda activate minirbt_intent\npip install torch transformers pandas scikit-learn onnx onnxruntime",
            "takeaway": "核心金句：环境隔离是工程落地素养的第一指标；干净、极简的工具链能帮你规避 80% 莫名其妙的环境玄学。"
        },
        {
            "kicker": "模块五 · 文本分析与小模型微调",
            "title": "微调 Step 1：分层抽样与 Tokenizer 张量化",
            "bullets": [
                "必须分层抽样（Stratified Split）：由于标签分布不均（优惠类仅占 10%），随机切分会导致验证集类别失真，必须按比例分层分流。",
                "数值张量化（Tokenization）：通过 MiniRBT 词表将中文文本转为离散整型 ID 序列（input_ids），首尾补入 [CLS] 与 [SEP] 标记。",
                "注意力掩码（attention_mask）：真实字符位设为 1 参与注意力计算，末端填充位（[PAD]）设为 0 强制屏蔽，阻断无用信息干扰。",
                "微批次理解：单机训练等效 Batch Size = DataLoader 批次大小 × 梯度累积步数（Gradient Accumulation），有效平滑显存压力。"
            ],
            "code": "# 严格分层抽样与定长张量化对齐\ntrain_df, val_df = train_test_split(df, test_size=0.2, stratify=df['label'])\ntk = AutoTokenizer.from_pretrained('hfl/minirbt-h256')\nbatch = tk(texts, padding='max_length', truncation=True, max_length=64, return_tensors='pt')",
            "takeaway": "核心金句：微调的第一步不是写训练代码，而是把字符串稳定转换成符合数学规范的张量流水线。"
        },
        {
            "kicker": "模块五 · 文本分析与小模型微调",
            "title": "微调 Step 2：前向传播与交叉熵损失（Loss）",
            "bullets": [
                "轻量编码底座：MiniRBT-H256 隐层维度仅 256（仅为标准 BERT 1/3），兼顾极快的响应速度与轻量参数体积，专精分类任务。",
                "分类头（Linear Head）：提取融合整句双向注意力的 [CLS] 语义特征向量，通过线性全连接层投射为 4 个意图类别的原始得分（Logits）。",
                "交叉熵度量（Cross-Entropy）：通过 Softmax 归一化为概率分布，损失函数单向重罚真实类别对应的负对数似然（-log P）。",
                "前向本质：让网络对当前样本做一次意图猜想，损失标量负责定量评价预测与真实业务标签之间的偏差程度。"
            ],
            "code": "# 加载预训练模型并在顶层挂载 4 分类线性头\nmodel = AutoModelForSequenceClassification.from_pretrained(\n    'hfl/minirbt-h256', num_labels=4\n)\nloss = model(**batch).loss  # 自动计算交叉熵损失",
            "takeaway": "核心金句：前向传播就是模型先做一次猜想，系统再用数学损失函数告诉它这次猜得有多离谱。"
        },
        {
            "kicker": "模块五 · 文本分析与小模型微调",
            "title": "微调 Step 3 & 4：责任溯源与 AdamW 更新",
            "bullets": [
                "反向传播的本质（Autograd）：标量 Loss 产生后，PyTorch 计算图通过微积分链式法则反向逆推，算出每个参数权重的偏导梯度（grad）。",
                "梯度物理含义：梯度揭示了“若将该参数稍微调大，全局损失是增加还是降低”，是模型参数寻找优化方向的罗盘指南。",
                "稳健更新（AdamW）：结合动量平滑与自适应学习率调节，并显式解耦权重衰减（Weight Decay），有效预防过拟合。",
                "必须清零梯度：PyTorch 设计上默认将反向传播的梯度持续累加（+=），每更新完一步必须显式调用 zero_grad() 清除历史梯度。"
            ],
            "code": "# 经典的训练四步循环范式\noptimizer.zero_grad()           # 1. 清空历史累积梯度\nloss = model(**batch).loss      # 2. 前向传播计算 Loss\nloss.backward()                 # 3. 反向传播链式求导\noptimizer.step()                # 4. AdamW 更新权重参数",
            "takeaway": "核心金句：反向传播是论功行过、明确每个参数该背多少责任；优化器是指引所有参数稳妥变聪明。"
        },

        # --- 模块六：模型轻量化与工程部署 ---
        {
            "kicker": "模块六 · 轻量化与部署",
            "title": "训练态与推理态的本质鸿沟",
            "bullets": [
                "生产环境的核心诉求：线上服务关注毫秒级低延迟（P99 Latency）、极少内存常驻、零依赖冲突与高并发稳定性。",
                "裸跑 PyTorch Serving 的缺陷：PyTorch 镜像通常数个 GB，携带冗余的自动求导引擎、优化器缓存与 Python GIL 锁性能消耗。",
                "状态彻底剥离：线上推理只保留纯前向路径（Forward Only），永久冻结 Dropout 与 BatchNorm，彻底关闭梯度计算（no_grad）。",
                "目标转变：从不断接收反馈、迭代优化的“学习态”，转换为输入张量直达预测输出的“纯静态确定性数学计算图”。"
            ],
            "code": "# 切换至推理模式并永久关闭梯度计算\nmodel.eval()\nwith torch.no_grad():\n    logits = model(**inputs).logits",
            "takeaway": "核心金句：训练是在学校里反复记录笔记和错题（梯度与优化器）；部署是上战场打仗，必须卸下书包轻装上阵。"
        },
        {
            "kicker": "模块六 · 轻量化与部署",
            "title": "跨平台图交换：ONNX 计算图固化",
            "bullets": [
                "开放标准规范：ONNX（Open Neural Network Exchange）定义了语言与硬件无关的静态计算图标准拓扑与底层算子集协议。",
                "图级物理优化：导出后推理引擎可在启动时完成算子融合（如 Gemm + Bias + GELU 合并单 Kernel 执行）与常量静态折叠。",
                "动态轴配置（Dynamic Axes）：生产环境文本长度与并发批次动态波动，导出时必须将 batch_size 与 seq_len 声明为动态变量。",
                "破除误区：ONNX 本身并不是压缩算法，它的核心使命是把动态 Python 代码描摹固化为脱离深度学习框架的静态计算拓扑。"
            ],
            "code": "# 导出带动态轴的通用 ONNX 静态计算图\ntorch.onnx.export(\n    model.eval(), (dummy['input_ids'], dummy['attention_mask']),\n    'artifacts/model.onnx',\n    input_names=['input_ids', 'attention_mask'],\n    output_names=['logits'],\n    dynamic_axes={'input_ids': {0: 'batch', 1: 'seq'}},\n    opset_version=17\n)",
            "takeaway": "核心金句：ONNX 就像数据科学界的 PDF 格式；不管当初用什么工具画的，导成 ONNX，任何生产引擎都能高效解读。"
        },
        {
            "kicker": "模块六 · 轻量化与部署",
            "title": "极限压缩：INT8 动态量化原理",
            "bullets": [
                "位宽降级数学原理：深度学习默认使用 32 位单精度浮点（FP32），每个参数占 4 字节；INT8 动态量化将其线性映射为 1 字节整型（-128~127）。",
                "体积收缩 75%：一个约 100MB 的模型权重经 INT8 动态量化后直接骤降至约 25MB，显著优化云镜像传输带宽与边缘设备加载内存。",
                "CPU 级指令集赋能：现代 CPU 具备专门的整数向量并行加速单元（如 VNNI、AVX-512 INT8），整型计算推理延迟显著低于浮点。",
                "工程权衡（Trade-off）：截断映射不可避免引入舍入噪声；上线前必须在测试集上校验量化前后的 F1-Score 精度损耗（通常 <0.5%）。"
            ],
            "code": "from onnxruntime.quantization import quantize_dynamic, QuantType\n# 一行代码执行工业级 INT8 动态量化\nquantize_dynamic(\n    'artifacts/model.onnx', 'artifacts/model.int8.onnx',\n    weight_type=QuantType.QInt8\n)",
            "takeaway": "核心金句：量化是工程艺术的极致；用 1% 的精度容忍代价，换取体积缩减至 1/4、吞吐倍增的巨大工程红利。"
        },
        {
            "kicker": "模块六 · 轻量化与部署",
            "title": "生产部署：纯 ONNX Runtime 轻量引擎",
            "bullets": [
                "斩断 Torch 依赖：生产 Serving 容器内无需安装庞大的 PyTorch，仅依赖极小体积的 onnxruntime C++ 绑定引擎与 NumPy 即可运行。",
                "纯粹三步流水线：文本通过轻量 Tokenizer 查表生成原生 ndarray $\\to$ 灌入 Session.run 极速推理 $\\to$ np.argmax 映射回中文意图。",
                "性能飞跃：消除 Python 解释器在动态构图上的开销，单次意图推理在普通 CPU 上可稳定压至数毫秒内，轻松应对高并发打点。",
                "生产资产交付物：生产目录最终仅需收敛为 model.int8.onnx 模型文件、分词词表与 label_mapping.json 标签映射字典。"
            ],
            "code": "# 生产级纯 CPU 无 Torch 依赖的高并发推理链路\nsession = ort.InferenceSession('artifacts/model.int8.onnx')\ninputs = {k: np.array(v) for k, v in encoded.items()}\nlogits = session.run(['logits'], inputs)[0]\npred_intent = id2label[int(np.argmax(logits, axis=-1)[0])]",
            "takeaway": "核心金句：生产系统越薄越健壮；没有 torch，没有 backward，只有输入数组与输出意图——这就是最强悍的工业落地姿态。"
        },

        # --- 模块七：全景收敛 ---
        {
            "kicker": "模块七 · 课程全景收敛",
            "title": "技术全景大闭环：从连续内存到业务变现",
            "bullets": [
                "NumPy（算力引擎）：通过数据缓冲区、步长系统与 SIMD 向量化彻底打破原生解释器低效，构筑同构张量底层性能基石。",
                "Pandas（业务中枢）：在矩阵上叠加 Index 标签对齐、缺失值守卫与 GroupBy、Merge 关系代数，赋予数据严肃的业务分析语义。",
                "微调训练（语义突破）：借助 MiniRBT 紧凑编码底座与前向反向优化闭环，攻克客服工单与用户非结构化意图分类难题。",
                "轻量部署（生产落地）：跨平台 ONNX 静态计算图固化结合 INT8 动态量化，以极简轻装完成工业级高并发毫秒级交付。"
            ],
            "code": "# 极简业务全景调用一览\ndf.query('amt > 0').groupby('hour')['amt'].transform('mean')\nquantize_dynamic('m.onnx', 'm.int8.onnx')\nsession.run(['logits'], ort_inputs)",
            "takeaway": "核心金句：Python 是当今时代的胶水，而我们掌握的是这瓶胶水里最硬核的骨架；带上对底层内存的敬畏去征服生产系统！"
        }
    ]

    for data in slides_data:
        slide = prs.slides.add_slide(blank_layout)

        # 1. 强制纯白空白背景
        bg = slide.background
        fill = bg.fill
        fill.solid()
        fill.fore_color.rgb = COLOR_BG

        # 2. 顶部模块指示标签（Kicker / Category）
        kicker_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.45), Inches(11.7), Inches(0.3))
        tf_k = kicker_box.text_frame
        tf_k.word_wrap = True
        tf_k.margin_left = tf_k.margin_right = tf_k.margin_top = tf_k.margin_bottom = 0
        pk = tf_k.paragraphs[0]
        pk.text = data["kicker"].upper()
        pk.font.name = FONT_MAIN
        pk.font.size = Pt(11)
        pk.font.bold = True
        pk.font.color.rgb = COLOR_SUBTITLE

        # 3. 主标题
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.75), Inches(11.7), Inches(0.6))
        tf_t = title_box.text_frame
        tf_t.word_wrap = True
        tf_t.margin_left = tf_t.margin_right = tf_t.margin_top = tf_t.margin_bottom = 0
        pt = tf_t.paragraphs[0]
        pt.text = data["title"]
        pt.font.name = FONT_MAIN
        pt.font.size = Pt(22)
        pt.font.bold = True
        pt.font.color.rgb = COLOR_TITLE

        # 4. 左侧：高密度原理与业务文字解释（字多、深挖细节）
        left_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.6), Inches(6.8), Inches(5.3))
        tf_l = left_box.text_frame
        tf_l.word_wrap = True
        tf_l.margin_left = tf_l.margin_right = tf_l.margin_top = tf_l.margin_bottom = 0

        for i, b_text in enumerate(data["bullets"]):
            p = tf_l.paragraphs[0] if i == 0 else tf_l.add_paragraph()
            p.text = f"•  {b_text}"
            p.font.name = FONT_MAIN
            p.font.size = Pt(12)
            p.font.color.rgb = COLOR_TEXT
            p.line_spacing = 1.25
            p.space_after = Pt(10)

        # 5. 右侧上方：极简核心代码展示卡片
        code_bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(7.9), Inches(1.6), Inches(4.6), Inches(3.2))
        code_bg.fill.solid()
        code_bg.fill.fore_color.rgb = COLOR_BOX_BG
        code_bg.line.color.rgb = COLOR_BOX_BORDER
        code_bg.line.width = Pt(1)

        code_box = slide.shapes.add_textbox(Inches(8.1), Inches(1.75), Inches(4.2), Inches(2.9))
        tf_c = code_box.text_frame
        tf_c.word_wrap = True
        tf_c.margin_left = tf_c.margin_right = tf_c.margin_top = tf_c.margin_bottom = 0

        # 代码标签
        p_ch = tf_c.paragraphs[0]
        p_ch.text = "CORE IMPLEMENTATION"
        p_ch.font.name = FONT_CODE
        p_ch.font.size = Pt(9)
        p_ch.font.bold = True
        p_ch.font.color.rgb = RGBColor(148, 163, 184)
        p_ch.space_after = Pt(8)

        # 代码本体
        p_code = tf_c.add_paragraph()
        p_code.text = data["code"]
        p_code.font.name = FONT_CODE
        p_code.font.size = Pt(11)
        p_code.font.color.rgb = COLOR_CODE
        p_code.line_spacing = 1.2

        # 6. 右侧下方：讲授重点 / 业务洞察金句卡片
        call_bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(7.9), Inches(5.05), Inches(4.6), Inches(1.85))
        call_bg.fill.solid()
        call_bg.fill.fore_color.rgb = COLOR_CALLOUT_BG
        call_bg.line.color.rgb = COLOR_CALLOUT_BORDER
        call_bg.line.width = Pt(1)

        call_box = slide.shapes.add_textbox(Inches(8.1), Inches(5.2), Inches(4.2), Inches(1.55))
        tf_cal = call_box.text_frame
        tf_cal.word_wrap = True
        tf_cal.margin_left = tf_cal.margin_right = tf_cal.margin_top = tf_cal.margin_bottom = 0

        p_calh = tf_cal.paragraphs[0]
        p_calh.text = "KEY TAKEAWAY & INSIGHT"
        p_calh.font.name = FONT_MAIN
        p_calh.font.size = Pt(9)
        p_calh.font.bold = True
        p_calh.font.color.rgb = RGBColor(79, 70, 229)
        p_calh.space_after = Pt(6)

        p_cal = tf_cal.add_paragraph()
        p_cal.text = data["takeaway"]
        p_cal.font.name = FONT_MAIN
        p_cal.font.size = Pt(11)
        p_cal.font.bold = True
        p_cal.font.color.rgb = COLOR_CALLOUT_TEXT
        p_cal.line_spacing = 1.2

    prs.save(output_path)
    print(f"PPT 生成完毕: {output_path} (共 {len(slides_data)} 页)")

if __name__ == "__main__":
    create_presentation()